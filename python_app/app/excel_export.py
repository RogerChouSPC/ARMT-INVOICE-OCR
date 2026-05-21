"""Write a list of InvoiceRow to an Excel file shaped like data/Output.xlsx."""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from app.schema import OUTPUT_COLUMNS, InvoiceRow


def write_excel(rows: list[InvoiceRow], path: str | Path | None = None) -> bytes:
    """Write `rows` to an .xlsx file.

    When `path` is None, returns the bytes (for FastAPI responses).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Output"

    ws.append(OUTPUT_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for i, row in enumerate(rows, start=1):
        ws.append(row.to_excel_row(i))

    # Reasonable column widths so the file is browsable.
    widths = [6, 32, 50, 16, 18, 14, 16, 16, 12, 22, 13, 13, 32, 60, 12, 10, 10, 10, 10, 10, 13, 28]
    for col_idx, w in enumerate(widths[: ws.max_column], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = w

    if path is not None:
        wb.save(path)
        return Path(path).read_bytes()
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
