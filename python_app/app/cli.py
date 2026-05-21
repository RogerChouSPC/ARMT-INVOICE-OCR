"""Command-line validator: run the pipeline on sample_invoices/ and diff vs Output.xlsx.

Usage from repo root:
    python -m app.cli                    # validate against the bundled gold
    python -m app.cli path/to/Invoice.pdf  # process one file, print rows

The validation report groups results by vendor and shows which gold rows
match exactly, which differ (and where), and which gold rows have no
corresponding extracted row.
"""

from __future__ import annotations

import argparse
import io
import sys
from pathlib import Path

import openpyxl

from app.customer_master import DEFAULT_MASTER_PATH, load_customer_master
from app.excel_export import write_excel
from app.extract import extract_pdf
from app.schema import InvoiceRow


REPO_ROOT = Path(__file__).resolve().parents[2]
SAMPLES_DIR = REPO_ROOT / "ARMT-INVOICE-OCR" / "data" / "sample_invoices"
GOLD_XLSX = REPO_ROOT / "ARMT-INVOICE-OCR" / "data" / "Output.xlsx"


# Columns we compare against Output.xlsx.
# Skipped on purpose:
#   - seq (positional)
#   - description / product_description (subject to OCR whitespace noise)
#   - remark / divisionsale (always blank in the gold)
#   - customergroup / customercode  ← excluded per user request: these come from
#     the Customer Master lookup, not from the invoice itself, so we judge the
#     OCR/parser accuracy against fields actually read from the PDF.
COMPARE_FIELDS = [
    ("taxid", "taxid", str),
    ("vendor_customercode", "vendor_customercode", str),
    ("vendor_branch", "vendor_branch", str),
    ("invoiceno", "invoiceno", str),
    ("amount", "amount", float),
    ("VAT 7% ", "vat_7", float),
    ("TAX % ", "tax_pct", float),
    ("TAX 2% ", "tax_2", float),
    ("TAX 3% ", "tax_3", float),
    ("TAX 5% ", "tax_5", float),
    ("netamount", "netamount", float),
]


def _f(v) -> float:
    if v in (None, "", 0):
        return 0.0
    try:
        return float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        return 0.0


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _load_gold() -> list[dict]:
    wb = openpyxl.load_workbook(GOLD_XLSX, data_only=True)
    ws = wb["Output"]
    hdr = [c.value for c in ws[1]]
    rows = []
    for r in range(2, ws.max_row + 1):
        row = {hdr[i]: ws.cell(r, i + 1).value for i in range(len(hdr))}
        rows.append(row)
    return rows


def _row_to_compare(r: InvoiceRow) -> dict:
    return {
        "customergroup": r.customergroup,
        "customercode": r.customercode,
        "taxid": r.taxid,
        "vendor_customercode": r.vendor_customercode,
        "vendor_branch": r.vendor_branch,
        "invoiceno": r.invoiceno,
        "amount": r.amount,
        "vat_7": r.vat_7,
        "tax_pct": r.tax_pct,
        "tax_2": r.tax_2,
        "tax_3": r.tax_3,
        "tax_5": r.tax_5,
        "netamount": r.netamount,
    }


def _key(d: dict) -> tuple[str, str]:
    """Match key = (taxid, invoiceno) — uniquely identifies a gold row (mostly)."""
    return (_s(d.get("taxid")), _s(d.get("invoiceno")))


def cmd_validate() -> int:
    master = load_customer_master(DEFAULT_MASTER_PATH)
    gold_rows = _load_gold()
    gold_by_key: dict[tuple[str, str], list[dict]] = {}
    for g in gold_rows:
        gold_by_key.setdefault(_key(g), []).append(g)

    # Process every PDF in samples folder
    extracted: list[tuple[str, InvoiceRow]] = []
    by_vendor: dict[str, int] = {}
    print(f"Scanning {SAMPLES_DIR} …\n")
    for pdf in sorted(SAMPLES_DIR.glob("*.pdf")):
        result = extract_pdf(pdf.read_bytes(), pdf.name, master)
        by_vendor[result.vendor_id or "(unknown)"] = (
            by_vendor.get(result.vendor_id or "(unknown)", 0) + len(result.rows)
        )
        status = "OCR" if result.used_ocr else "TXT"
        warn = f"  ⚠ {result.warnings[0]}" if result.warnings else ""
        print(f"  [{status}] {pdf.name:30} vendor={result.vendor_id or '?':10} pages={result.page_count} rows={len(result.rows)}{warn}")
        for r in result.rows:
            extracted.append((pdf.name, r))

    print(f"\nExtracted {len(extracted)} row(s) total. By vendor:")
    for v, n in sorted(by_vendor.items()):
        print(f"  {v:12} {n}")

    # Compare against gold (customer master fields excluded per user request).
    print(f"\n--- Validation vs Output.xlsx ({len(gold_rows)} gold rows) ---")
    print("    Compared fields:", ", ".join(o for _, o, _ in COMPARE_FIELDS))
    print()

    # Group gold by vendor taxid → human-readable label from the customergroup.
    def vendor_label(g: dict) -> str:
        return _s(g.get("customergroup")).split(" - ", 1)[0] or _s(g.get("taxid"))

    per_vendor: dict[str, dict[str, int]] = {}
    diff_details: list[str] = []

    # Group extracted rows by (taxid, invoiceno). Multi-line invoices (e.g. the
    # LT credit note has 7 line items under one invoice number, CP ALL has 22)
    # will have many entries per key.
    extracted_index: dict[tuple[str, str], list[InvoiceRow]] = {}
    for _, r in extracted:
        extracted_index.setdefault((r.taxid, r.invoiceno), []).append(r)

    # Greedy best-match: for each gold row, pick the not-yet-consumed extracted
    # row under the same key with the closest `amount`. This stops all gold
    # rows of the same invoice from collapsing onto candidates[0].
    consumed: dict[tuple[str, str], set[int]] = {key: set() for key in extracted_index}

    for gold in gold_rows:
        label = vendor_label(gold)
        stats = per_vendor.setdefault(label, {"matched": 0, "differ": 0, "missing": 0, "gold": 0})
        stats["gold"] += 1

        key = _key(gold)
        candidates = extracted_index.get(key, [])
        if not candidates:
            stats["missing"] += 1
            continue

        # Pick the unused extracted row with the closest amount.
        g_amount = _f(gold.get("amount"))
        used = consumed[key]
        best_idx = -1
        best_diff = float("inf")
        for i, c in enumerate(candidates):
            if i in used:
                continue
            d = abs(g_amount - _f(c.amount))
            if d < best_diff:
                best_diff = d
                best_idx = i
        if best_idx < 0:
            # All candidates consumed by earlier gold rows — fall back to the
            # first one so we still produce a diff report.
            best_idx = 0
        else:
            used.add(best_idx)

        ours = _row_to_compare(candidates[best_idx])
        diffs = []
        for gold_col, ours_col, kind in COMPARE_FIELDS:
            g_val = gold.get(gold_col)
            o_val = ours.get(ours_col)
            if kind is float:
                if abs(_f(g_val) - _f(o_val)) > 0.01:
                    diffs.append(f"{ours_col}: gold={_f(g_val)} ours={_f(o_val)}")
            else:
                if _s(g_val) != _s(o_val):
                    diffs.append(f"{ours_col}: gold={_s(g_val)!r} ours={_s(o_val)!r}")
        if diffs:
            stats["differ"] += 1
            diff_details.append(f"  ≠ [{label}] {key[1]} (amt {g_amount}): {diffs[:3]}")
        else:
            stats["matched"] += 1

    # Per-vendor table
    print(f"  {'vendor':<32} {'gold':>5} {'match':>6} {'diff':>5} {'miss':>5}")
    print(f"  {'-'*32} {'-'*5} {'-'*6} {'-'*5} {'-'*5}")
    total = {"matched": 0, "differ": 0, "missing": 0, "gold": 0}
    for label in sorted(per_vendor):
        s = per_vendor[label]
        status = "✓" if s["differ"] == 0 and s["missing"] == 0 else (
            "≠" if s["differ"] > 0 else "-"
        )
        print(f"  {status} {label:<30} {s['gold']:>5} {s['matched']:>6} {s['differ']:>5} {s['missing']:>5}")
        for k in total: total[k] += s[k]
    print(f"  {'-'*32} {'-'*5} {'-'*6} {'-'*5} {'-'*5}")
    print(f"  {'TOTAL':<32} {total['gold']:>5} {total['matched']:>6} {total['differ']:>5} {total['missing']:>5}")

    if diff_details:
        print("\nDifferences:")
        for d in diff_details:
            print(d)
    print(
        f"\n  Vendors with extracted rows: "
        f"{', '.join(v for v, n in sorted(by_vendor.items()) if n > 0) or '(none)'}"
    )
    print(
        "  Vendors with 0 extracted rows are the ones whose parser is not yet implemented (see warnings above)."
    )
    return 0 if total["differ"] == 0 else 1


def cmd_one(path: Path) -> int:
    master = load_customer_master(DEFAULT_MASTER_PATH)
    result = extract_pdf(path.read_bytes(), path.name, master)
    print(f"Vendor: {result.vendor_id}, pages: {result.page_count}, rows: {len(result.rows)}, OCR: {result.used_ocr}")
    for w in result.warnings:
        print(f"  ⚠ {w}")
    for r in result.rows:
        print()
        for k, v in r.to_dict().items():
            if v not in (None, "", 0, 0.0):
                print(f"  {k}: {v!r}")
    # Also drop a quick Excel beside the PDF
    out_xlsx = path.with_suffix(".extracted.xlsx")
    write_excel(result.rows, out_xlsx)
    print(f"\nWrote {out_xlsx}")
    return 0


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="ARMT invoice extractor")
    p.add_argument("path", nargs="?", help="A single PDF to process; omit to run validation suite")
    args = p.parse_args(argv)
    # Force UTF-8 stdout on Windows so Thai prints correctly
    if sys.stdout.encoding and "utf" not in sys.stdout.encoding.lower():
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    if args.path:
        return cmd_one(Path(args.path))
    return cmd_validate()


if __name__ == "__main__":
    raise SystemExit(main())
